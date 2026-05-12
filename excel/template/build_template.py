"""
Gera FutebolLocal.xlsx a partir do zero.

Não faz parte da aplicação Symfony — é a ferramenta usada para (re)gerar o
template sempre que a estrutura das abas mudar, em vez de editar o .xlsx
manualmente e perder o controle de versão do que mudou.

Uso: pip install openpyxl && python build_template.py
(rodar a partir desta pasta — o arquivo é escrito em ./FutebolLocal.xlsx)
"""
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------- paleta (mesma do design system do projeto) ----------
BG = "101110"
SURFACE = "171917"
LINE = "2A2E2A"
TEXT = "F0F2EC"
MUTED = "A4AAA1"
ACCENT = "C5F05C"
ACCENT_DARK = "1B2410"

header_font = Font(name="Calibri", bold=True, color=TEXT, size=11)
header_fill = PatternFill("solid", fgColor=ACCENT_DARK)
title_font = Font(name="Calibri", bold=True, color=ACCENT, size=16)
subtitle_font = Font(name="Calibri", color=MUTED, size=10)
label_font = Font(name="Calibri", bold=True, color=TEXT, size=10)
thin = Side(style="thin", color=LINE)
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()
wb.remove(wb.active)

DATA_ROWS = 200  # linhas em branco disponíveis em cada tabela


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(vertical="center")


def add_table(ws, name, first_row, headers, extra_rows=DATA_ROWS):
    ncols = len(headers)
    for i, h in enumerate(headers, start=1):
        ws.cell(row=first_row, column=i, value=h)
    style_header_row(ws, first_row, ncols)
    last_row = first_row + extra_rows
    last_col = get_column_letter(ncols)
    ref = f"A{first_row}:{last_col}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
    )
    ws.add_table(table)
    for c in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = max(14, len(headers[c - 1]) + 4)
    ws.freeze_panes = ws.cell(row=first_row + 1, column=1)
    return last_row


def add_validation(ws, dv, cell_range):
    ws.add_data_validation(dv)
    dv.add(cell_range)


# ================= LISTAS (listas de apoio p/ validação, aba oculta) =================
ws_listas = wb.create_sheet("LISTAS")
listas = {
    "A": ("Status geral", ["ATIVO", "INATIVO"]),
    "B": ("Posição", ["Goleiro", "Zagueiro", "Lateral", "Meio-campo", "Atacante"]),
    "C": ("Status do jogo", ["SCHEDULED", "LIVE", "FINISHED", "POSTPONED", "CANCELLED"]),
    "D": ("Tipo de evento", ["GOAL", "ASSIST", "YELLOW_CARD", "RED_CARD", "SUBSTITUTION_IN", "SUBSTITUTION_OUT", "OWN_GOAL"]),
    "E": ("Titular", ["SIM", "NAO"]),
}
for col, (label, values) in listas.items():
    ws_listas[f"{col}1"] = label
    ws_listas[f"{col}1"].font = header_font
    for i, v in enumerate(values, start=2):
        ws_listas[f"{col}{i}"] = v
ws_listas.sheet_state = "hidden"

def list_formula(col, count):
    return f"=LISTAS!${col}$2:${col}${count + 1}"


# ================= README =================
ws = wb.create_sheet("README")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 100
ws["A1"] = "Futebol Local — Planilha de sincronização"
ws["A1"].font = title_font
ws["A2"] = "Use esta planilha para cadastrar e atualizar dados do campeonato offline. O VBA envia os dados para a API do Futebol Local, que valida, processa e persiste."
ws["A2"].font = subtitle_font
ws["A2"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[2].height = 30

instructions = [
    ("Abas", ""),
    ("CONFIG", "Preencha a URL da API e o token antes de sincronizar."),
    ("TIMES / JOGADORES / ELENCOS / RODADAS / JOGOS / EVENTOS / ESCALACOES", "Cadastre os dados nas tabelas. Colunas com lista suspensa só aceitam os valores válidos."),
    ("Logo (URL) em TIMES", "Opcional — cole a URL de uma imagem já hospedada em algum lugar. Para enviar um arquivo de imagem direto do computador, use o painel administrativo (Times > editar time) em vez da planilha."),
    ("", ""),
    ("Como sincronizar", ""),
    ("1.", "Preencha os dados nas abas."),
    ("2.", "Na aba CONFIG, clique em \"Validar planilha\" para checar erros antes de enviar."),
    ("3.", "Clique em \"Sincronizar\" para enviar os dados para a API."),
    ("4.", "Acompanhe o status e a data da última sincronização na própria aba CONFIG."),
    ("", ""),
    ("Habilitando as macros (VBA)", ""),
    ("1.", "Salve este arquivo como .xlsm (Pasta de Trabalho Habilitada para Macro)."),
    ("2.", "Abra o Editor VBA (Alt+F11) e importe os módulos de excel/vba/*.bas."),
    ("3.", "Na aba CONFIG, insira os três botões (Desenvolvedor > Inserir > Botão) e associe às macros SincronizarDados, ValidarPlanilha e BaixarDados."),
    ("", "Veja excel/documentation/vba-setup.md para o passo a passo completo."),
]
r = 4
for a, b in instructions:
    ws.cell(row=r, column=1, value=a).font = label_font if b == "" and a not in ("1.", "2.", "3.", "4.", "") else Font(name="Calibri", size=10, color=TEXT)
    if b:
        ws.cell(row=r, column=1, value=f"{a}  {b}" if a in ("1.", "2.", "3.", "4.") else a)
        if a not in ("1.", "2.", "3.", "4."):
            ws.cell(row=r, column=2, value=b).font = Font(name="Calibri", size=10, color=MUTED)
    r += 1

# ================= CONFIG =================
ws = wb.create_sheet("CONFIG")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 46
ws.column_dimensions["D"].width = 22
ws["A1"] = "Configuração"
ws["A1"].font = title_font

rows = [
    ("Organização (slug)", "liga-municipal-januaria"),
    ("Campeonato (slug)", "campeonato-municipal-2026"),
    ("Temporada (ID)", ""),
    ("URL da API", "http://127.0.0.1:8000/api/v1"),
    ("Token da API", ""),
]
named_cells = {}
r = 3
for label, default in rows:
    ws.cell(row=r, column=1, value=label).font = label_font
    cell = ws.cell(row=r, column=2, value=default)
    cell.border = border
    cell.fill = PatternFill("solid", fgColor=SURFACE)
    cell.font = Font(name="Calibri", color=TEXT)
    named_cells[label] = (r, cell)
    r += 1

r += 1
ws.cell(row=r, column=1, value="Status").font = label_font
status_cell = ws.cell(row=r, column=2, value="Não sincronizado")
status_cell.border = border
status_cell.fill = PatternFill("solid", fgColor=SURFACE)
status_row = r
r += 1
ws.cell(row=r, column=1, value="Última sincronização").font = label_font
last_sync_cell = ws.cell(row=r, column=2, value="")
last_sync_cell.border = border
last_sync_cell.fill = PatternFill("solid", fgColor=SURFACE)
last_sync_row = r

# nomes definidos, para o VBA referenciar por nome em vez de endereço fixo
defined = {
    "cfgOrgSlug": f"CONFIG!$B$3",
    "cfgChampionshipSlug": f"CONFIG!$B$4",
    "cfgSeasonId": f"CONFIG!$B$5",
    "cfgApiUrl": f"CONFIG!$B$6",
    "cfgApiToken": f"CONFIG!$B$7",
    "cfgStatus": f"CONFIG!$B${status_row}",
    "cfgLastSync": f"CONFIG!$B${last_sync_row}",
}
for name, ref in defined.items():
    wb.defined_names[name] = openpyxl.workbook.defined_name.DefinedName(name, attr_text=ref)

r += 3
ws.cell(row=r, column=1, value="Ações (após habilitar as macros — ver aba README)").font = label_font
r += 1
for label in ["[ Baixar dados ]", "[ Validar planilha ]", "[ Sincronizar ]"]:
    cell = ws.cell(row=r, column=1, value=label)
    cell.font = Font(name="Calibri", bold=True, color=BG)
    cell.fill = PatternFill("solid", fgColor=ACCENT)
    cell.alignment = Alignment(horizontal="center")
    cell.border = border
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1

# ================= TIMES =================
ws = wb.create_sheet("TIMES")
last = add_table(ws, "TabelaTimes", 1, ["ID", "Nome", "Nome curto", "Cidade", "Status", "Logo (URL)"])
add_validation(ws, DataValidation(type="list", formula1=list_formula("A", 2), allow_blank=True), f"E2:E{last}")
ws["A2"] = 1; ws["B2"] = "União FC"; ws["C2"] = "UNI"; ws["D2"] = "Januária"; ws["E2"] = "ATIVO"
ws["A3"] = 2; ws["B3"] = "Real América"; ws["C3"] = "REA"; ws["D3"] = "Januária"; ws["E3"] = "ATIVO"

# ================= JOGADORES =================
ws = wb.create_sheet("JOGADORES")
last = add_table(ws, "TabelaJogadores", 1, ["ID", "Nome", "Data nascimento", "Posição", "Número", "Status"])
add_validation(ws, DataValidation(type="list", formula1=list_formula("B", 5), allow_blank=True), f"D2:D{last}")
add_validation(ws, DataValidation(type="list", formula1=list_formula("A", 2), allow_blank=True), f"F2:F{last}")
ws["A2"] = 1; ws["B2"] = "João Silva"; ws["C2"] = "2000-05-10"; ws["D2"] = "Atacante"; ws["E2"] = 9; ws["F2"] = "ATIVO"

# ================= ELENCOS =================
ws = wb.create_sheet("ELENCOS")
last = add_table(ws, "TabelaElencos", 1, ["ID", "Time", "Jogador", "Temporada", "Número", "Data início", "Data fim", "Status"])
add_validation(ws, DataValidation(type="list", formula1=list_formula("A", 2), allow_blank=True), f"H2:H{last}")
ws["A2"] = 1; ws["B2"] = "União FC"; ws["C2"] = "João Silva"; ws["D2"] = "2026"; ws["E2"] = 9; ws["H2"] = "ATIVO"

# ================= RODADAS =================
ws = wb.create_sheet("RODADAS")
last = add_table(ws, "TabelaRodadas", 1, ["ID", "Número", "Nome", "Data", "Status"])
add_validation(ws, DataValidation(type="list", formula1=list_formula("A", 2), allow_blank=True), f"E2:E{last}")
ws["A2"] = 1; ws["B2"] = 1; ws["C2"] = "Rodada 1"; ws["D2"] = "2026-09-20"; ws["E2"] = "ATIVO"

# ================= JOGOS =================
ws = wb.create_sheet("JOGOS")
last = add_table(ws, "TabelaJogos", 1, ["ID", "Rodada", "Data", "Hora", "Mandante", "Visitante", "Estádio", "Placar Mandante", "Placar Visitante", "Status"])
add_validation(ws, DataValidation(type="list", formula1=list_formula("C", 5), allow_blank=True), f"J2:J{last}")
ws["A2"] = 1; ws["B2"] = 1; ws["C2"] = "2026-09-20"; ws["D2"] = "16:00"; ws["E2"] = "União FC"; ws["F2"] = "Real América"; ws["G2"] = "Estádio Municipal"; ws["J2"] = "SCHEDULED"

# ================= EVENTOS =================
ws = wb.create_sheet("EVENTOS")
last = add_table(ws, "TabelaEventos", 1, ["ID", "Jogo", "Minuto", "Time", "Jogador", "Tipo"])
add_validation(ws, DataValidation(type="list", formula1=list_formula("D", 7), allow_blank=True), f"F2:F{last}")
ws["A2"] = 1; ws["B2"] = 1; ws["C2"] = 23; ws["D2"] = "União FC"; ws["E2"] = "João Silva"; ws["F2"] = "GOAL"

# ================= ESCALACOES =================
ws = wb.create_sheet("ESCALACOES")
last = add_table(ws, "TabelaEscalacoes", 1, ["ID", "Jogo", "Time", "Jogador", "Titular", "Posição", "Número", "Entrada", "Saída"])
add_validation(ws, DataValidation(type="list", formula1=list_formula("E", 2), allow_blank=True), f"E2:E{last}")
add_validation(ws, DataValidation(type="list", formula1=list_formula("B", 5), allow_blank=True), f"F2:F{last}")
ws["A2"] = 1; ws["B2"] = 1; ws["C2"] = "União FC"; ws["D2"] = "João Silva"; ws["E2"] = "SIM"; ws["F2"] = "Atacante"; ws["G2"] = 9

# ordena as abas
order = ["README", "CONFIG", "TIMES", "JOGADORES", "ELENCOS", "RODADAS", "JOGOS", "EVENTOS", "ESCALACOES", "LISTAS"]
wb._sheets = [wb[name] for name in order]
wb.active = 0

import os
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "FutebolLocal.xlsx")
wb.save(output_path)
print(f"OK: {output_path} gerado")
